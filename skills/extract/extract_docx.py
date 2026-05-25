from __future__ import annotations

from pathlib import Path

from skills.extract.extract_txt import TextExtractor, TextExtractorError


class DocxExtractorError(TextExtractorError):
    pass


class DocxExtractor(TextExtractor):
    def _read_document_text(self, source_path: Path) -> str:
        try:
            from docx import Document
        except ModuleNotFoundError as exc:
            raise DocxExtractorError(
                "python-docx is not installed. Please run `python -m pip install -r requirements-dev.txt`."
            ) from exc

        try:
            document = Document(str(source_path))
        except Exception as exc:
            raise DocxExtractorError(f"DOCX document in {source_path} could not be read.") from exc

        text_parts: list[str] = []
        text_parts.extend(
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text and paragraph.text.strip()
        )
        for table in document.tables:
            for row in table.rows:
                cell_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text and cell.text.strip())
                if cell_text:
                    text_parts.append(cell_text)

        if not text_parts:
            raise DocxExtractorError(f"DOCX document in {source_path} did not contain extractable text.")

        return "\n\n".join(text_parts)
